import streamlit as st
from onboarding import generate_onboarding_plan
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

st.set_page_config(page_title="Onboarding Assistant", page_icon="🌸", layout="centered")

# ---------- State ----------
defaults = {
    "step": 1,
    "job_title": "Project Manager",
    "department": "Project Management",
    "experience_level": "Junior",
    "tasks": None,
    "completed_tasks": {},
}
for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value

# ---------- Styling ----------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Playfair+Display:wght@400;500;600;700&display=swap');
html, body, [class*="css"] {font-family:'DM Sans',sans-serif;color:#423340;}
h1,h2,h3 {font-family:'Playfair Display',serif!important;color:#423340;font-weight:500;}
.stApp {background:linear-gradient(145deg,#FFF7F8 0%,#FFF9FC 55%,#F8F3FC 100%);}
.block-container {max-width:1100px;padding-top:2.5rem;padding-bottom:3rem;}
div[data-testid="stAlert"] {border-radius:16px;border:1px solid #F2D5DF;}
div[data-testid="stExpander"] {background:rgba(255,255,255,.88);border:1px solid #F0DCE4;border-radius:16px;overflow:hidden;}
div[data-testid="stProgressBar"]>div {background:#F3E5EB;border-radius:999px;}
div[data-testid="stProgressBar"]>div>div {background:linear-gradient(90deg,#F2B8C9,#E9CDEB);border-radius:999px;}
div.stButton>button,div.stDownloadButton>button {
 background:linear-gradient(135deg,#F6D5E0,#EFB9CA);color:#493744;border:1px solid #EAB4C7;
 border-radius:16px;font-family:'DM Sans',sans-serif;font-size:1rem;font-weight:700;padding:.75rem 1.2rem;
 box-shadow:0 5px 14px rgba(194,132,158,.16);transition:all .2s ease;}
div.stButton>button:hover,div.stDownloadButton>button:hover {
 background:linear-gradient(135deg,#F0C2D2,#E7A5BC);color:#493744;border-color:#E7A5BC;
 box-shadow:0 8px 20px rgba(194,132,158,.25);transform:translateY(-2px);}
div[data-testid="stSelectbox"]>div>div {background:#FFF;border-radius:12px;border-color:#EBDDE3;}
.hero {background:linear-gradient(135deg,#FBE8EE,#F7F0FA);border:1px solid #F0DCE5;border-radius:24px;padding:2.5rem;margin-bottom:1.8rem;}
.hero-eyebrow {color:#A36F89;font-size:1.1rem;font-weight:700;letter-spacing:.12em;margin-bottom:1rem;}
.hero h1 {font-family:'DM Sans',sans-serif!important;font-size:clamp(2.2rem,6vw,3.6rem);line-height:1.1;letter-spacing:-.04em;margin:0 0 1rem;color:#493744;font-weight:700;}
.hero p {color:#796775;font-size:1.05rem;margin:0;line-height:1.7;}
.section-heading {margin:2rem 0 1.5rem;}
.section-label {display:inline-block;color:#A36F89;background:#F9E5EC;border-radius:999px;padding:.45rem 1rem;font-size:.78rem;font-weight:700;letter-spacing:.12em;}
.section-heading h2 {font-family:'DM Sans',sans-serif!important;font-size:clamp(1.7rem,4vw,2.5rem);font-weight:700;color:#493744;margin:.8rem 0 .4rem;}
.section-heading p {color:#796775;font-size:1rem;margin:0;}
.step-caption {color:#8B7484;font-size:.85rem;text-align:right;margin-bottom:.35rem;}
.step-track {display:flex;gap:.5rem;margin:.5rem 0 1.65rem;}
.step-segment {height:6px;flex:1;border-radius:999px;background:#EDE3EB;}
.step-segment.active {background:linear-gradient(90deg,#E8AFC3,#D9BDE8);}
.step-names {display:flex;gap:.5rem;margin-top:-1.2rem;margin-bottom:2rem;}
.step-name {flex:1;font-size:.75rem;color:#A3919E;}
.step-name.active {color:#A36F89;font-weight:700;}
.priority-badge {display:inline-block;padding:.3rem .75rem;border-radius:999px;font-size:.78rem;font-weight:700;letter-spacing:.02em;}
.priority-high {background:#F7D4DF;color:#8F3F60;border:1px solid #EDB8CA;}
.priority-medium {background:#FCEAF0;color:#98647B;border:1px solid #F2D7E2;}
.priority-low {background:#EEE7F8;color:#765B98;border:1px solid #DDD0EF;}
.info-card {background:rgba(255,255,255,.72);border:1px solid #F0DCE5;border-radius:16px;padding:1rem 1.2rem;}
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="hero">
 <div class="hero-eyebrow">A structured beginning.</div>
 <h1>Onboarding Process Generator</h1>
 <p>Create a thoughtful onboarding experience, tailored to every new hire.</p>
</div>
""", unsafe_allow_html=True)


def render_progress(step):
    names = ["Employee details", "Review & generate", "Track & export"]
    segments, labels = [], []

    for i, name in enumerate(names, start=1):
        active = "active" if i <= step else ""
        segments.append(f'<div class="step-segment {active}"></div>')
        labels.append(f'<div class="step-name {active}">0{i} · {name}</div>')

    st.markdown(
        f'<div class="step-caption">STEP {step} OF 3</div>'
        f'<div class="step-track">{"".join(segments)}</div>'
        f'<div class="step-names">{"".join(labels)}</div>',
        unsafe_allow_html=True
    )


def heading(label, title, description):
    st.markdown(
        f'<div class="section-heading"><span class="section-label">{label}</span>'
        f'<h2>{title}</h2><p>{description}</p></div>',
        unsafe_allow_html=True
    )


render_progress(st.session_state.step)

# ---------- Step 1 ----------
if st.session_state.step == 1:
    heading(
        "STEP 01",
        "Let's meet your new hire",
        "A few details will help us create a more relevant onboarding plan."
    )

    c1, c2 = st.columns(2, gap="large")

    with c1:
        st.selectbox(
            "Job title",
            ["Project Manager", "BI Analyst", "Marketing Specialist"],
            key="job_title"
        )

    with c2:
        st.selectbox(
            "Department",
            ["Project Management", "Data & Analytics", "Marketing", "Operations"],
            key="department"
        )

    st.selectbox(
        "Experience level",
        ["Junior", "Mid-level", "Senior"],
        key="experience_level"
    )

    if st.button("Continue to review →", type="primary", use_container_width=True):
        st.session_state.tasks = None
        st.session_state.completed_tasks = {}
        st.session_state.step = 2
        st.rerun()

# ---------- Step 2 ----------
elif st.session_state.step == 2:
    heading(
        "STEP 02",
        "Review the details",
        "Make sure everything looks right before generating the onboarding plan."
    )

    st.markdown(
        f'<div class="info-card"><strong>Job title</strong><br>{st.session_state.job_title}<br><br>'
        f'<strong>Department</strong><br>{st.session_state.department}<br><br>'
        f'<strong>Experience level</strong><br>{st.session_state.experience_level}</div>',
        unsafe_allow_html=True
    )

    st.write("")
    back, generate = st.columns(2, gap="medium")

    with back:
        if st.button("← Back", use_container_width=True):
            st.session_state.step = 1
            st.rerun()

    with generate:
        if st.button("Generate my plan →", type="primary", use_container_width=True):
            st.session_state.tasks = generate_onboarding_plan(
                st.session_state.job_title,
                st.session_state.department,
                st.session_state.experience_level
            )
            st.session_state.completed_tasks = {
                i: False
                for i in range(1, len(st.session_state.tasks) + 1)
            }
            st.session_state.step = 3
            st.rerun()

# ---------- Step 3 ----------
elif st.session_state.step == 3:
    heading(
        "STEP 03",
        "Your onboarding plan",
        "Track progress, explore each task, and export your plan when you're ready."
    )

    tasks = st.session_state.tasks or []

    if not tasks:
        st.warning("There is no onboarding plan yet. Please generate one first.")

        if st.button("Back to details"):
            st.session_state.step = 1
            st.rerun()

        st.stop()

    st.write(f"**Job title:** {st.session_state.job_title}")
    st.write(f"**Department:** {st.session_state.department}")
    st.write(f"**Experience level:** {st.session_state.experience_level}")

    st.warning("Draft plan — manager review is recommended.")

total = len(tasks)

completed_count = sum(
    st.session_state.get(f"task_completed_{i}", False)
    for i in range(1, total + 1)
    )

progress = completed_count / total if total else 0

    st.subheader("Onboarding progress")
    st.progress(progress)
    st.write(f"{completed_count} of {total} tasks completed ({progress:.0%})")

    if progress == 1:
        st.success("Congratulations! Onboarding completed successfully!")
        st.balloons()

    st.divider()
    st.subheader("Tasks")

    for number, task in enumerate(tasks, start=1):
        completed = st.checkbox(
            f"Mark as completed — {task['task']}",
            key=f"task_completed_{number}",
            value=st.session_state.completed_tasks.get(number, False)
        )
        st.session_state.completed_tasks[number] = completed

        with st.expander(f"{number}. {task['task']}", expanded=False):
            st.write(f"**Phase:** {task['phase']}")

            priority = task["priority"].lower()
            priority_class = (
                "priority-high" if priority == "high"
                else "priority-medium" if priority == "medium"
                else "priority-low"
            )

            st.markdown(
                f'<span class="priority-badge {priority_class}">'
                f'Priority · {task["priority"]}</span>',
                unsafe_allow_html=True
            )

            st.write(f"**Owner:** {task['owner']}")
            st.write(
                f"**Independence level:** "
                f"{task.get('independence_level', 'Not specified')}"
            )
            st.write("**Instructions:**")

            for instruction in task["instructions"]:
                st.write(f"- {instruction}")

            st.write(f"**Expected outcome:** {task['expected_outcome']}")

    # ---------- TXT export ----------
    plan_text = "ONBOARDING PLAN\n" + "=" * 40 + "\n\n"
    plan_text += (
        f"Job title: {st.session_state.job_title}\n"
        f"Department: {st.session_state.department}\n"
        f"Experience level: {st.session_state.experience_level}\n\n"
    )

    for number, task in enumerate(tasks, start=1):
        status = (
            "Completed"
            if st.session_state.completed_tasks.get(number, False)
            else "Not started"
        )

        plan_text += (
            f"{number}. {task['task']}\n"
            f"   Phase: {task['phase']}\n"
            f"   Priority: {task['priority']}\n"
            f"   Owner: {task['owner']}\n"
            f"   Status: {status}\n"
            f"   Instructions:\n"
        )

        for instruction in task["instructions"]:
            plan_text += f"   - {instruction}\n"

        plan_text += f"   Expected outcome: {task['expected_outcome']}\n\n"

    # ---------- Excel export ----------
    wb = Workbook()
    ws = wb.active
    ws.title = "Onboarding Plan"

    ws.append([
        "Phase",
        "Task",
        "Priority",
        "Owner",
        "Instructions",
        "Expected Outcome",
        "Status"
    ])

    for number, task in enumerate(tasks, start=1):
        status = (
            "Completed"
            if st.session_state.completed_tasks.get(number, False)
            else "Not started"
        )

        ws.append([
            task["phase"],
            task["task"],
            task["priority"],
            task["owner"],
            "\n".join(task["instructions"]),
            task["expected_outcome"],
            status
        ])

    header_fill = PatternFill(fill_type="solid", fgColor="E9E1F5")

    for cell in ws[1]:
        cell.font = Font(bold=True, color="423340")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col_num, width in enumerate([16, 32, 14, 20, 48, 48, 16], start=1):
        ws.column_dimensions[get_column_letter(col_num)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    validation = DataValidation(
        type="list",
        formula1='"Not started,In progress,Completed"',
        allow_blank=False
    )
    validation.error = "Please select a status from the dropdown."
    validation.errorTitle = "Invalid status"
    validation.prompt = "Select the current task status."
    validation.promptTitle = "Task status"
    ws.add_data_validation(validation)

    for row_num in range(2, ws.max_row + 1):
        validation.add(f"G{row_num}")

    colors = {
        "Not started": ("FCE4D6", "9C0006"),
        "In progress": ("FFF2CC", "7F6000"),
        "Completed": ("E2F0D9", "375623")
    }

    for status, (bg, fg) in colors.items():
        ws.conditional_formatting.add(
            f"G2:G{ws.max_row}",
            CellIsRule(
                operator="equal",
                formula=[f'"{status}"'],
                fill=PatternFill(fill_type="solid", fgColor=bg),
                font=Font(color=fg)
            )
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    excel = BytesIO()
    wb.save(excel)
    excel.seek(0)

    st.divider()
    st.subheader("Export your onboarding plan")
    st.caption("Choose the format that works best for you.")

    txt_col, xlsx_col = st.columns(2, gap="medium")

    with txt_col:
        st.download_button(
            "📄 Download as TXT",
            data=plan_text,
            file_name="onboarding_plan.txt",
            mime="text/plain",
            use_container_width=True
        )

    with xlsx_col:
        st.download_button(
            "📊 Download as Excel (.xlsx)",
            data=excel.getvalue(),
            file_name="onboarding_plan.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

    if st.button("← Start a new onboarding plan", use_container_width=True):
        st.session_state.step = 1
        st.session_state.tasks = None
        st.session_state.completed_tasks = {}
        st.rerun()
